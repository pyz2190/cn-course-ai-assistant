"""真实文档解析器：支持 PDF、PPT、字幕文件的解析。

将课程文档解析为结构化的 ChunkMetadata，每个 Chunk 携带章节、页码和
知识点关联信息，保证检索片段可追溯到原始材料。
"""

from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from uuid import NAMESPACE_URL, uuid5

from app.domain.enums import (
    AccessLevel,
    ContentType,
    Language,
    ParseStatus,
)
from app.domain.models import ChunkMetadata


@dataclass
class ParsedChunk:
    """解析后的文档片段。"""

    content: str
    chapter: str
    page_start: int | None
    page_end: int | None
    title: str | None = None
    extracted_content_type: ContentType | None = None
    source_path: str | None = None


class DocumentParser:
    """真实文档解析器。

    支持解析 PDF、PPT、字幕文件，提取文本内容并按章节切分为 Chunk。
    """

    def parse(self, file_path: Path, content_type: ContentType) -> list[ParsedChunk]:
        """解析文档并返回切分后的 Chunk 列表。

        Args:
            file_path: 文档文件路径
            content_type: 文档类型（pdf/ppt/subtitle）

        Returns:
            解析后的 Chunk 列表

        Raises:
            ValueError: 不支持的文档类型
            FileNotFoundError: 文件不存在
        """
        if not file_path.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        if content_type == ContentType.PDF:
            return self._parse_pdf(file_path)
        elif content_type == ContentType.PPT:
            return self._parse_ppt(file_path)
        elif content_type == ContentType.SUBTITLE:
            return self._parse_subtitle(file_path)
        else:
            raise ValueError(f"不支持的文档类型: {content_type}")

    def _parse_pdf(self, file_path: Path) -> list[ParsedChunk]:
        """解析 PDF 文档。

        使用 PyMuPDF 提取文本、表格、图片和公式相关内容。
        按页面切分并识别章节结构。
        """
        import pymupdf

        chunks: list[ParsedChunk] = []
        image_dir = file_path.parent / f"{file_path.stem}_images"
        image_count = 0

        with pymupdf.open(str(file_path)) as doc:
            current_chapter = "前言"
            chapter_content: list[str] = []
            chapter_start_page = 1

            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text("text")

                # 提取表格
                table_chunks = self._extract_pdf_tables(
                    page, current_chapter, page_num + 1
                )
                chunks.extend(table_chunks)

                # 提取图片
                img_count, img_chunks = self._extract_pdf_images(
                    page, image_dir, image_count, current_chapter, page_num + 1
                )
                image_count += img_count
                chunks.extend(img_chunks)

                # 检测数学公式（基于 Unicode 数学符号）
                formula_lines = self._detect_formulas(text)
                if formula_lines:
                    chunks.append(
                        ParsedChunk(
                            content="\n".join(formula_lines),
                            chapter=current_chapter,
                            page_start=page_num + 1,
                            page_end=page_num + 1,
                            title=f"公式（第 {page_num + 1} 页）",
                            extracted_content_type=ContentType.FORMULA,
                        )
                    )

                # 章节识别 + 文本提取
                lines = text.split("\n")
                for line in lines:
                    line = line.strip()
                    if self._is_chapter_header(line):
                        if chapter_content:
                            chunks.append(
                                ParsedChunk(
                                    content="\n".join(chapter_content),
                                    chapter=current_chapter,
                                    page_start=chapter_start_page,
                                    page_end=page_num,
                                )
                            )
                            chapter_content = []
                        current_chapter = line
                        chapter_start_page = page_num + 1

                    if line and not self._is_header_footer(line):
                        chapter_content.append(line)

            # 保存最后一个章节
            if chapter_content:
                chunks.append(
                    ParsedChunk(
                        content="\n".join(chapter_content),
                        chapter=current_chapter,
                        page_start=chapter_start_page,
                        page_end=len(doc),
                    )
                )

        return chunks

    def _extract_pdf_tables(
        self, page, chapter: str, page_num: int
    ) -> list[ParsedChunk]:
        """从 PDF 页面提取表格。"""
        chunks: list[ParsedChunk] = []
        try:
            tables = page.find_tables()
            for idx, table in enumerate(tables):
                # 将表格转为 Markdown 格式
                rows = []
                for row in table.extract():
                    cells = [str(cell).strip() if cell else "" for cell in row]
                    rows.append("| " + " | ".join(cells) + " |")
                    # 第一行后加表头分隔线
                    if len(rows) == 1:
                        sep = "| " + " | ".join(["---"] * len(cells)) + " |"
                        rows.append(sep)

                if rows:
                    chunks.append(
                        ParsedChunk(
                            content="\n".join(rows),
                            chapter=chapter,
                            page_start=page_num,
                            page_end=page_num,
                            title=f"表格 {idx + 1}（第 {page_num} 页）",
                            extracted_content_type=ContentType.TABLE,
                        )
                    )
        except Exception:
            pass  # 表格提取失败不影响其他内容
        return chunks

    def _extract_pdf_images(
        self,
        page,
        image_dir: Path,
        start_index: int,
        chapter: str,
        page_num: int,
    ) -> tuple[int, list[ParsedChunk]]:
        """从 PDF 页面提取图片。返回 (新增图片数, chunks)。"""
        chunks: list[ParsedChunk] = []
        count = 0
        try:
            images = page.get_images(full=True)
            for img_info in images:
                xref = img_info[0]
                try:
                    pix = self._extract_image_from_doc(page.parent, xref)
                    if pix is None:
                        continue

                    image_dir.mkdir(parents=True, exist_ok=True)
                    img_path = image_dir / f"img_{start_index + count:04d}.png"
                    pix.save(str(img_path))
                    count += 1

                    chunks.append(
                        ParsedChunk(
                            content=f"[图片] 页面 {page_num} 的图片已提取到: {img_path.name}",
                            chapter=chapter,
                            page_start=page_num,
                            page_end=page_num,
                            title=f"图片（第 {page_num} 页）",
                            extracted_content_type=ContentType.IMAGE,
                            source_path=str(img_path),
                        )
                    )
                except Exception:
                    continue
        except Exception:
            pass
        return count, chunks

    @staticmethod
    def _extract_image_from_doc(doc, xref: int):
        """从 PDF 文档中提取指定 xref 的图片。"""
        import pymupdf

        try:
            return pymupdf.Pixmap(doc, xref)
        except Exception:
            return None

    @staticmethod
    def _detect_formulas(text: str) -> list[str]:
        """检测文本中的数学公式行。

        基于 Unicode 数学符号和常见公式模式进行检测。
        """
        import re

        math_patterns = [
            r"[∑∏∫∂∇√∞±×÷≈≠≤≥∈∉⊂⊃∪∩]",
            r"\b[a-zA-Z]\s*[=]\s*[\d.]+\s*[+\-*/]",
            r"\b[sin|cos|tan|log|ln|exp]\s*\(",
            r"\b\d+\s*[+\-*/]\s*\d+\s*=",
            r"[αβγδεζηθικλμνξπρστυφχψω]",
        ]
        combined = re.compile("|".join(math_patterns))

        formula_lines = []
        for line in text.split("\n"):
            line = line.strip()
            if line and combined.search(line):
                formula_lines.append(line)
        return formula_lines

    def _parse_ppt(self, file_path: Path) -> list[ParsedChunk]:
        """解析 PPT 文档。

        使用 python-pptx 提取幻灯片内容，包括文本、表格和图片。
        """
        from pptx import Presentation

        chunks: list[ParsedChunk] = []
        prs = Presentation(str(file_path))
        image_dir = file_path.parent / f"{file_path.stem}_images"
        image_count = 0

        for slide_num, slide in enumerate(prs.slides, start=1):
            title = ""
            content_lines: list[str] = []
            has_diagram = False

            for shape in slide.shapes:
                # 提取表格
                if shape.has_table:
                    table = shape.table
                    rows = []
                    for row in table.rows:
                        cells = [
                            cell.text.strip() if cell.text else ""
                            for cell in row.cells
                        ]
                        rows.append("| " + " | ".join(cells) + " |")
                        if len(rows) == 1:
                            sep = "| " + " | ".join(["---"] * len(cells)) + " |"
                            rows.append(sep)
                    if rows:
                        chapter = title if title else f"幻灯片 {slide_num}"
                        chunks.append(
                            ParsedChunk(
                                content="\n".join(rows),
                                chapter=chapter,
                                page_start=slide_num,
                                page_end=slide_num,
                                title=f"表格（幻灯片 {slide_num}）",
                                extracted_content_type=ContentType.TABLE,
                            )
                        )

                # 提取文本
                elif shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            if not title and shape.shape_type == 13:
                                title = text
                            else:
                                content_lines.append(text)

                # 检测图片
                elif hasattr(shape, "image"):
                    try:
                        image_dir.mkdir(parents=True, exist_ok=True)
                        img_blob = shape.image.blob
                        ext = shape.image.content_type.split("/")[-1]
                        img_path = image_dir / f"slide_{slide_num}_img_{image_count}.{ext}"
                        img_path.write_bytes(img_blob)
                        image_count += 1
                        content_lines.append(f"[图片] 已提取: {img_path.name}")
                    except Exception:
                        pass

                # 检测组合形状（可能是时序图/流程图）
                if hasattr(shape, "shapes") and len(getattr(shape, "shapes", [])) > 2:
                    has_diagram = True

            # 如果检测到可能是时序图/流程图的组合形状
            if has_diagram:
                chapter = title if title else f"幻灯片 {slide_num}"
                chunks.append(
                    ParsedChunk(
                        content=f"[图表] 幻灯片 {slide_num} 包含组合图形（可能是时序图或流程图）",
                        chapter=chapter,
                        page_start=slide_num,
                        page_end=slide_num,
                        title=f"图表（幻灯片 {slide_num}）",
                        extracted_content_type=ContentType.DIAGRAM,
                    )
                )

            if content_lines:
                chapter = title if title else f"幻灯片 {slide_num}"
                chunks.append(
                    ParsedChunk(
                        content="\n".join(content_lines),
                        chapter=chapter,
                        page_start=slide_num,
                        page_end=slide_num,
                        title=title if title else None,
                    )
                )

        return chunks

    def _parse_subtitle(self, file_path: Path) -> list[ParsedChunk]:
        """解析字幕文件。

        支持 SRT 和 WebVTT 格式，按时间窗口切分为 Chunk。
        """
        content = file_path.read_text(encoding="utf-8")
        suffix = file_path.suffix.lower()

        if suffix == ".srt":
            entries = self._parse_srt(content)
        elif suffix == ".vtt":
            entries = self._parse_vtt(content)
        else:
            raise ValueError(f"不支持的字幕格式: {suffix}")

        # 按时间窗口切分（每5分钟一个 Chunk）
        chunks: list[ParsedChunk] = []
        window_size = 300  # 5分钟
        current_window_start = 0
        current_content: list[str] = []

        for entry in entries:
            if entry.start_time >= current_window_start + window_size:
                # 保存当前窗口
                if current_content:
                    window_end = current_window_start + window_size
                    time_range = f"{_fmt_time(current_window_start)}-{_fmt_time(window_end)}"
                    idx = len(chunks) + 1
                    chunks.append(
                        ParsedChunk(
                            content="\n".join(current_content),
                            chapter=f"字幕片段 {idx} [{time_range}]",
                            page_start=None,
                            page_end=None,
                        )
                    )
                    current_content = []

                current_window_start = entry.start_time

            current_content.append(entry.text)

        # 保存最后一个窗口
        if current_content:
            window_end = current_window_start + window_size
            time_range = f"{_fmt_time(current_window_start)}-{_fmt_time(window_end)}"
            idx = len(chunks) + 1
            chunks.append(
                ParsedChunk(
                    content="\n".join(current_content),
                    chapter=f"字幕片段 {idx} [{time_range}]",
                    page_start=None,
                    page_end=None,
                )
            )

        return chunks

    def _parse_srt(self, content: str) -> list["_SubtitleEntry"]:
        """解析 SRT 格式字幕。"""
        entries: list[_SubtitleEntry] = []
        blocks = content.strip().split("\n\n")

        for block in blocks:
            lines = block.strip().split("\n")
            if len(lines) >= 3:
                # 解析时间戳
                time_line = lines[1]
                start_str, end_str = time_line.split(" --> ")
                start_time = self._srt_time_to_seconds(start_str.strip())
                end_time = self._srt_time_to_seconds(end_str.strip())
                text = "\n".join(lines[2:])

                entries.append(_SubtitleEntry(start_time, end_time, text))

        return entries

    def _parse_vtt(self, content: str) -> list["_SubtitleEntry"]:
        """解析 WebVTT 格式字幕。"""
        entries: list[_SubtitleEntry] = []
        lines = content.strip().split("\n")

        # 跳过 WEBVTT 头部，找到第一个时间戳行
        i = 0
        while i < len(lines) and "-->" not in lines[i]:
            i += 1

        while i < len(lines):
            if "-->" in lines[i]:
                time_line = lines[i]
                start_str, end_str = time_line.split(" --> ")
                start_time = self._vtt_time_to_seconds(start_str.strip())
                end_time = self._vtt_time_to_seconds(end_str.strip())

                # 收集文本行
                text_lines: list[str] = []
                i += 1
                while i < len(lines) and lines[i].strip():
                    text_lines.append(lines[i].strip())
                    i += 1

                if text_lines:
                    entries.append(
                        _SubtitleEntry(start_time, end_time, "\n".join(text_lines))
                    )
            else:
                i += 1

        return entries

    def _srt_time_to_seconds(self, time_str: str) -> float:
        """将 SRT 时间格式转换为秒数。"""
        # 格式: HH:MM:SS,mmm
        time_str = time_str.replace(",", ".")
        parts = time_str.split(":")
        hours = int(parts[0])
        minutes = int(parts[1])
        seconds = float(parts[2])
        return hours * 3600 + minutes * 60 + seconds

    def _vtt_time_to_seconds(self, time_str: str) -> float:
        """将 VTT 时间格式转换为秒数。"""
        # 格式: HH:MM:SS.mmm
        parts = time_str.split(":")
        hours = int(parts[0])
        minutes = int(parts[1])
        seconds = float(parts[2])
        return hours * 3600 + minutes * 60 + seconds

    def _is_chapter_header(self, line: str) -> bool:
        """检测是否为章节标题。"""
        import re

        # 中文章节：第X章、第X节、第 X 章、第 X 节
        if re.match(r"^第\s*[一二三四五六七八九十\d]+\s*[章节]", line):
            return True
        # 英文章节：Chapter X、Section X
        return bool(re.match(r"^(Chapter|Section)\s+\d+", line, re.IGNORECASE))

    def _is_header_footer(self, line: str) -> bool:
        """检测是否为页眉页脚。"""
        return len(line) < 3 or line.isdigit()


@dataclass
class _SubtitleEntry:
    """字幕条目。"""

    start_time: float
    end_time: float
    text: str


def _fmt_time(seconds: float) -> str:
    """将秒数格式化为 HH:MM:SS。"""
    h = int(seconds // 3600)
    m = int((seconds % 3600) // 60)
    s = int(seconds % 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


class RealResourceImporter:
    """B 模块：真实课程资料导入适配器。

    解析 PDF、PPT、字幕文件，生成带元数据的 Chunk，每个 Chunk 携带章节、页码和
    知识点关联信息，保证检索片段可追溯到原始材料。
    """

    def __init__(self, storage_dir: Path | None = None) -> None:
        self._parser = DocumentParser()
        self._storage_dir = storage_dir or Path("data/resources")
        self._storage_dir.mkdir(parents=True, exist_ok=True)

    def import_resource(
        self,
        file_path: Path,
        course_id: str,
        title: str,
        version: str,
        language: Language,
        content_type: ContentType,
        knowledge_point_ids: list[str] | None = None,
    ) -> list[ChunkMetadata]:
        """导入课程资料并生成 Chunk Metadata。

        Args:
            file_path: 文档文件路径
            course_id: 课程 ID
            title: 资料标题
            version: 版本号
            language: 语言
            content_type: 内容类型
            knowledge_point_ids: 知识点 ID 列表（可选）

        Returns:
            生成的 ChunkMetadata 列表
        """
        # 生成资源 ID
        resource_key = f"{course_id}:{title}:{version}"
        resource_id = f"resource-{uuid5(NAMESPACE_URL, resource_key).hex[:12]}"

        # 解析文档
        parsed_chunks = self._parser.parse(file_path, content_type)

        # 生成 ChunkMetadata
        chunks: list[ChunkMetadata] = []
        for i, parsed in enumerate(parsed_chunks):
            chunk_id = f"chunk-{resource_id}-{i:03d}"
            chunk = ChunkMetadata(
                chunk_id=chunk_id,
                resource_id=resource_id,
                knowledge_point_ids=knowledge_point_ids or ["kp-pending-review"],
                title=parsed.title or title,
                content=parsed.content,
                chapter=parsed.chapter,
                page_start=parsed.page_start,
                page_end=parsed.page_end,
                language=language,
                content_type=content_type,
                extracted_content_type=parsed.extracted_content_type,
                source_url=None,
                source_path=parsed.source_path,
                version=version,
                access_level=AccessLevel.COURSE,
                parse_status=ParseStatus.PARSED,
                updated_at=datetime.now(UTC),
            )
            chunks.append(chunk)

        return chunks

    def import_all_chunks(
        self,
        file_path: Path,
        course_id: str,
        title: str,
        version: str,
        language: Language,
        content_type: ContentType,
        knowledge_point_ids: list[str] | None = None,
    ) -> list[ChunkMetadata]:
        """导入一份资料的全部 Chunk，用于批量构建知识库。

        这是 import_resource 的别名，保持接口一致性。
        """
        return self.import_resource(
            file_path=file_path,
            course_id=course_id,
            title=title,
            version=version,
            language=language,
            content_type=content_type,
            knowledge_point_ids=knowledge_point_ids,
        )

    def import_from_file(
        self,
        file_path: Path,
        course_id: str,
        title: str,
        version: str,
        language: str,
        content_type: str,
        knowledge_point_ids: list[str] | None = None,
    ) -> list[ChunkMetadata]:
        """从文件路径导入，接受字符串类型的 language 和 content_type。

        适配 FileResourceImporter 协议，方便路由层直接调用。
        """
        return self.import_resource(
            file_path=file_path,
            course_id=course_id,
            title=title,
            version=version,
            language=Language(language),
            content_type=ContentType(content_type),
            knowledge_point_ids=knowledge_point_ids,
        )
