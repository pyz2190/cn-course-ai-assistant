"""B 模块测试：真实文档解析器。"""

from pathlib import Path

import pytest

from app.adapters.document_parser import DocumentParser, RealResourceImporter
from app.domain.enums import ContentType, Language


@pytest.fixture
def parser():
    return DocumentParser()


@pytest.fixture
def importer():
    return RealResourceImporter()


@pytest.fixture
def sample_pdf(tmp_path):
    """创建一个简单的 PDF 测试文件。"""
    import pymupdf

    pdf_path = tmp_path / "test.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "第 1 章 计算机网络概述", fontname="china-s")
    page.insert_text((72, 100), "这是第一章的内容。", fontname="china-s")
    page = doc.new_page()
    page.insert_text((72, 72), "第 2 章 应用层", fontname="china-s")
    page.insert_text((72, 100), "这是第二章的内容。", fontname="china-s")
    doc.save(str(pdf_path))
    doc.close()
    return pdf_path


@pytest.fixture
def sample_ppt(tmp_path):
    """创建一个简单的 PPT 测试文件。"""
    from pptx import Presentation

    pptx_path = tmp_path / "test.pptx"
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]  # 标题幻灯片

    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "计算机网络概述"

    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "应用层协议"

    prs.save(str(pptx_path))
    return pptx_path


@pytest.fixture
def sample_srt(tmp_path):
    """创建一个简单的 SRT 字幕测试文件。"""
    content = """1
00:00:01,000 --> 00:00:05,000
欢迎来到计算机网络课程

2
00:00:05,500 --> 00:00:10,000
今天我们将学习TCP协议

3
00:00:10,500 --> 00:00:15,000
TCP是传输层的重要协议
"""
    srt_path = tmp_path / "test.srt"
    srt_path.write_text(content, encoding="utf-8")
    return srt_path


class TestDocumentParser:
    """DocumentParser 单元测试。"""

    def test_parse_pdf_returns_chunks(self, parser, sample_pdf):
        chunks = parser.parse(sample_pdf, ContentType.PDF)
        assert len(chunks) > 0
        assert all(chunk.content for chunk in chunks)
        assert all(chunk.chapter for chunk in chunks)

    def test_parse_pdf_identifies_chapters(self, parser, sample_pdf):
        chunks = parser.parse(sample_pdf, ContentType.PDF)
        chapters = [chunk.chapter for chunk in chunks]
        assert any("第 1 章" in ch for ch in chapters)
        assert any("第 2 章" in ch for ch in chapters)

    def test_parse_ppt_returns_chunks(self, parser, sample_ppt):
        chunks = parser.parse(sample_ppt, ContentType.PPT)
        assert len(chunks) > 0
        assert all(chunk.content for chunk in chunks)

    def test_parse_subtitle_returns_chunks(self, parser, sample_srt):
        chunks = parser.parse(sample_srt, ContentType.SUBTITLE)
        assert len(chunks) > 0
        assert all(chunk.content for chunk in chunks)

    def test_parse_nonexistent_file_raises_error(self, parser):
        with pytest.raises(FileNotFoundError):
            parser.parse(Path("nonexistent.pdf"), ContentType.PDF)

    def test_parse_unsupported_type_raises_error(self, parser, sample_pdf):
        with pytest.raises(ValueError):
            parser.parse(sample_pdf, ContentType.TEXT)


class TestRealResourceImporter:
    """RealResourceImporter 单元测试。"""

    def test_import_resource_returns_chunks(self, importer, sample_pdf):
        chunks = importer.import_resource(
            file_path=sample_pdf,
            course_id="course-cn-2026",
            title="计算机网络：自顶向下方法",
            version="8e-zh",
            language=Language.ZH,
            content_type=ContentType.PDF,
        )
        assert len(chunks) > 0
        assert all(chunk.resource_id for chunk in chunks)
        assert all(chunk.chunk_id for chunk in chunks)
        assert all(chunk.title == "计算机网络：自顶向下方法" for chunk in chunks)

    def test_import_resource_has_metadata(self, importer, sample_pdf):
        chunks = importer.import_resource(
            file_path=sample_pdf,
            course_id="course-cn-2026",
            title="测试资料",
            version="v1",
            language=Language.ZH,
            content_type=ContentType.PDF,
        )
        chunk = chunks[0]
        assert chunk.language == Language.ZH
        assert chunk.content_type == ContentType.PDF
        assert chunk.version == "v1"
        assert chunk.access_level == "course"
        assert chunk.parse_status == "parsed"

    def test_import_resource_with_knowledge_points(self, importer, sample_pdf):
        chunks = importer.import_resource(
            file_path=sample_pdf,
            course_id="course-cn-2026",
            title="测试资料",
            version="v1",
            language=Language.ZH,
            content_type=ContentType.PDF,
            knowledge_point_ids=["kp-network-overview"],
        )
        assert all(chunk.knowledge_point_ids == ["kp-network-overview"] for chunk in chunks)
