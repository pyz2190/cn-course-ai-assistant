"""文件上传端点测试：覆盖 PDF、PPT、字幕文件的上传解析流程。"""

import io

from fastapi.testclient import TestClient


class TestUploadEndpoint:
    """POST /api/v1/resources/upload 测试。"""

    def _make_pdf_bytes(self) -> bytes:
        """生成一个简单的 PDF 文件字节流。"""
        import pymupdf

        doc = pymupdf.open()
        page = doc.new_page()
        page.insert_text((72, 72), "第 1 章 计算机网络概述", fontname="china-s")
        page.insert_text((72, 100), "这是测试内容。", fontname="china-s")
        pdf_bytes = doc.tobytes()
        doc.close()
        return pdf_bytes

    def _make_pptx_bytes(self) -> bytes:
        """生成一个简单的 PPTX 文件字节流。"""
        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "测试幻灯片"
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()

    def _make_srt_bytes(self) -> bytes:
        """生成一个简单的 SRT 字幕文件字节流。"""
        content = """1
00:00:01,000 --> 00:00:05,000
欢迎来到计算机网络课程

2
00:00:05,500 --> 00:00:10,000
今天学习TCP协议
"""
        return content.encode("utf-8")

    def test_upload_pdf_success(self, client: TestClient) -> None:
        pdf_bytes = self._make_pdf_bytes()
        response = client.post(
            "/api/v1/resources/upload",
            files={"file": ("test.pdf", pdf_bytes, "application/pdf")},
            data={
                "course_id": "course-cn-2026",
                "title": "测试教材",
                "version": "v1",
                "language": "zh",
                "content_type": "pdf",
                "knowledge_point_ids": "kp-network-overview",
            },
        )
        assert response.status_code == 201
        payload = response.json()
        assert payload["filename"] == "test.pdf"
        assert payload["chunk_count"] > 0
        assert payload["parse_status"] == "parsed"
        assert payload["resource_id"].startswith("resource-")
        assert payload["error"] is None
        assert len(payload["chunks"]) > 0

    def test_upload_pptx_success(self, client: TestClient) -> None:
        pptx_bytes = self._make_pptx_bytes()
        response = client.post(
            "/api/v1/resources/upload",
            files={"file": ("test.pptx", pptx_bytes, "application/vnd.ms-powerpoint")},
            data={
                "course_id": "course-cn-2026",
                "title": "测试课件",
                "version": "v1",
                "language": "zh",
                "content_type": "ppt",
            },
        )
        assert response.status_code == 201
        payload = response.json()
        assert payload["filename"] == "test.pptx"
        assert payload["parse_status"] == "parsed"

    def test_upload_srt_success(self, client: TestClient) -> None:
        srt_bytes = self._make_srt_bytes()
        response = client.post(
            "/api/v1/resources/upload",
            files={"file": ("test.srt", srt_bytes, "application/x-subrip")},
            data={
                "course_id": "course-cn-2026",
                "title": "测试字幕",
                "version": "v1",
                "language": "zh",
                "content_type": "subtitle",
            },
        )
        assert response.status_code == 201
        payload = response.json()
        assert payload["filename"] == "test.srt"
        assert payload["parse_status"] == "parsed"
        assert payload["chunk_count"] > 0

    def test_upload_vtt_success(self, client: TestClient) -> None:
        vtt_content = """WEBVTT

00:00:01.000 --> 00:00:05.000
Welcome to computer networks

00:00:05.500 --> 00:00:10.000
Today we learn TCP
"""
        response = client.post(
            "/api/v1/resources/upload",
            files={"file": ("test.vtt", vtt_content.encode("utf-8"), "text/vtt")},
            data={
                "course_id": "course-cn-2026",
                "title": "测试VTT字幕",
                "version": "v1",
                "language": "en",
                "content_type": "subtitle",
            },
        )
        assert response.status_code == 201
        payload = response.json()
        assert payload["filename"] == "test.vtt"
        assert payload["parse_status"] == "parsed"

    def test_upload_invalid_content_type(self, client: TestClient) -> None:
        response = client.post(
            "/api/v1/resources/upload",
            files={"file": ("test.pdf", b"fake", "application/pdf")},
            data={
                "course_id": "course-cn-2026",
                "title": "测试",
                "version": "v1",
                "language": "zh",
                "content_type": "invalid_type",
            },
        )
        assert response.status_code == 422
        assert "不支持的 content_type" in response.json()["message"]

    def test_upload_invalid_language(self, client: TestClient) -> None:
        pdf_bytes = self._make_pdf_bytes()
        response = client.post(
            "/api/v1/resources/upload",
            files={"file": ("test.pdf", pdf_bytes, "application/pdf")},
            data={
                "course_id": "course-cn-2026",
                "title": "测试",
                "version": "v1",
                "language": "fr",
                "content_type": "pdf",
            },
        )
        assert response.status_code == 422
        assert "不支持的 language" in response.json()["message"]

    def test_upload_extension_mismatch(self, client: TestClient) -> None:
        """文件扩展名与 content_type 不匹配应返回错误。"""
        response = client.post(
            "/api/v1/resources/upload",
            files={"file": ("test.txt", b"some content", "text/plain")},
            data={
                "course_id": "course-cn-2026",
                "title": "测试",
                "version": "v1",
                "language": "zh",
                "content_type": "pdf",
            },
        )
        assert response.status_code == 422
        assert "不匹配" in response.json()["message"]

    def test_upload_ppt_mismatch_extension(self, client: TestClient) -> None:
        """PPT content_type 但文件是 .txt 应报错。"""
        response = client.post(
            "/api/v1/resources/upload",
            files={"file": ("test.txt", b"some content", "text/plain")},
            data={
                "course_id": "course-cn-2026",
                "title": "测试",
                "version": "v1",
                "language": "zh",
                "content_type": "ppt",
            },
        )
        assert response.status_code == 422

    def test_upload_chunks_have_correct_metadata(self, client: TestClient) -> None:
        """验证上传后 Chunk 的元数据完整性。"""
        pdf_bytes = self._make_pdf_bytes()
        response = client.post(
            "/api/v1/resources/upload",
            files={"file": ("test.pdf", pdf_bytes, "application/pdf")},
            data={
                "course_id": "course-cn-2026",
                "title": "元数据测试",
                "version": "v2",
                "language": "zh",
                "content_type": "pdf",
                "knowledge_point_ids": "kp-test-001,kp-test-002",
            },
        )
        assert response.status_code == 201
        payload = response.json()
        chunk = payload["chunks"][0]
        assert chunk["resource_id"] == payload["resource_id"]
        assert chunk["title"] == "元数据测试"
        assert chunk["version"] == "v2"
        assert chunk["language"] == "zh"
        assert chunk["content_type"] == "pdf"
        assert chunk["parse_status"] == "parsed"
        assert chunk["access_level"] == "course"
        assert "kp-test-001" in chunk["knowledge_point_ids"]
        assert "kp-test-002" in chunk["knowledge_point_ids"]

    def test_upload_multiple_knowledge_points(self, client: TestClient) -> None:
        """验证多个知识点 ID 用逗号分隔传入。"""
        pdf_bytes = self._make_pdf_bytes()
        response = client.post(
            "/api/v1/resources/upload",
            files={"file": ("test.pdf", pdf_bytes, "application/pdf")},
            data={
                "course_id": "course-cn-2026",
                "title": "多知识点测试",
                "version": "v1",
                "language": "zh",
                "content_type": "pdf",
                "knowledge_point_ids": "kp-a, kp-b, kp-c",
            },
        )
        assert response.status_code == 201
        chunk = response.json()["chunks"][0]
        assert chunk["knowledge_point_ids"] == ["kp-a", "kp-b", "kp-c"]

    def test_upload_empty_knowledge_points_uses_default(self, client: TestClient) -> None:
        """不传知识点时使用默认值。"""
        pdf_bytes = self._make_pdf_bytes()
        response = client.post(
            "/api/v1/resources/upload",
            files={"file": ("test.pdf", pdf_bytes, "application/pdf")},
            data={
                "course_id": "course-cn-2026",
                "title": "默认知识点测试",
                "version": "v1",
                "language": "zh",
                "content_type": "pdf",
            },
        )
        assert response.status_code == 201
        chunk = response.json()["chunks"][0]
        assert chunk["knowledge_point_ids"] == ["kp-pending-review"]

    def test_upload_other_content_type_success(self, client: TestClient) -> None:
        """content_type=other 现在支持 .txt 文件上传。"""
        response = client.post(
            "/api/v1/resources/upload",
            files={"file": ("test.txt", b"some content here", "text/plain")},
            data={
                "course_id": "course-cn-2026",
                "title": "测试",
                "version": "v1",
                "language": "zh",
                "content_type": "other",
            },
        )
        assert response.status_code == 201
        assert response.json()["parse_status"] == "parsed"

    def test_upload_unsupported_extension_returns_422(self, client: TestClient) -> None:
        """不支持的文件扩展名应返回 422。"""
        response = client.post(
            "/api/v1/resources/upload",
            files={"file": ("test.xyz", b"content", "application/octet-stream")},
            data={
                "course_id": "course-cn-2026",
                "title": "测试",
                "version": "v1",
                "language": "zh",
                "content_type": "pdf",
            },
        )
        assert response.status_code == 422
        assert "不匹配" in response.json()["message"]

    def test_upload_rfc_content_type(self, client: TestClient) -> None:
        """RFC 类型应支持 .txt 文件。"""
        response = client.post(
            "/api/v1/resources/upload",
            files={"file": ("rfc791.txt", b"RFC 791 content", "text/plain")},
            data={
                "course_id": "course-cn-2026",
                "title": "RFC 791",
                "version": "v1",
                "language": "en",
                "content_type": "rfc",
            },
        )
        assert response.status_code == 201

    def test_upload_text_content_type(self, client: TestClient) -> None:
        """TEXT 类型应支持 .txt 和 .md 文件。"""
        response = client.post(
            "/api/v1/resources/upload",
            files={"file": ("notes.md", b"# Chapter 1\n\nThis is content.", "text/markdown")},
            data={
                "course_id": "course-cn-2026",
                "title": "课程笔记",
                "version": "v1",
                "language": "zh",
                "content_type": "text",
            },
        )
        assert response.status_code == 201
        assert response.json()["parse_status"] == "parsed"
        assert response.json()["chunk_count"] > 0
